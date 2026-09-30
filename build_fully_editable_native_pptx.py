"""
Build 100% Native, Fully Editable PowerPoint Deck for IPsight SIH 2026
Every single card, title, bullet, and flowchart box is a native PPT shape/textbox.
No flat images or uneditable overlays. Everything can be edited directly in PowerPoint.
"""

import os
import shutil
import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def build_fully_editable_deck():
    src_template = r"C:\Users\AS\Downloads\SIH2026-IDEA-Presentation-Format.pptx"
    output_pptx = r"C:\Users\AS\Downloads\vpn\IPsight_SIH26160_Submission.pptx"
    output_pdf = r"C:\Users\AS\Downloads\vpn\IPsight_SIH26160_Submission.pdf"
    mirror_pptx = r"C:\Users\AS\Downloads\SIH2026_IDEA_IPsight_Praxis.pptx"
    mirror_pdf = r"C:\Users\AS\Downloads\SIH2026_IDEA_IPsight_Praxis.pdf"

    # Close any running PowerPoint instances to avoid file lock
    os.system('powershell -Command "Get-Process POWERPNT -ErrorAction SilentlyContinue | Stop-Process -Force"')

    shutil.copyfile(src_template, output_pptx)
    prs = pptx.Presentation(output_pptx)

    # Color Palette - Professional Defense & Cyber Theme
    C_NAVY_DARK = RGBColor(11, 37, 69)       # #0B2545
    C_NAVY_TITLE = RGBColor(15, 23, 42)      # #0F172A
    C_BLUE_MED = RGBColor(30, 58, 138)       # #1E3A8A
    C_BLUE_ACCENT = RGBColor(2, 132, 199)    # #0284C7
    C_CYAN_BG = RGBColor(240, 249, 255)      # #F0F9FF
    C_LIGHT_BG = RGBColor(248, 250, 252)     # #F8FAFC
    C_CARD_BORDER = RGBColor(203, 213, 225)  # Slate 300
    C_WHITE = RGBColor(255, 255, 255)
    C_TEXT_DARK = RGBColor(15, 23, 42)       # Slate 900
    C_TEXT_BODY = RGBColor(51, 65, 85)       # Slate 700
    C_TEXT_MUTED = RGBColor(100, 116, 139)   # Slate 500
    C_GREEN = RGBColor(16, 149, 94)          # Emerald #10955E
    C_AMBER = RGBColor(217, 119, 6)          # Amber #D97706
    C_RED = RGBColor(220, 38, 38)            # Red #DC2626
    C_PURPLE = RGBColor(109, 40, 217)        # Violet #6D28D9

    def style_run(run, text, size=11.0, color=C_TEXT_BODY, bold=False, italic=False, font_name="Calibri"):
        run.text = text
        run.font.size = Pt(size)
        run.font.color.rgb = color
        run.font.bold = bold
        run.font.italic = italic
        run.font.name = font_name

    def clear_textbox(shape):
        tf = shape.text_frame
        tf.word_wrap = True
        p0 = tf.paragraphs[0]
        p0.text = ""
        p_elements = list(tf._txBody.p_lst)
        for p_elem in p_elements[1:]:
            tf._txBody.remove(p_elem)

    def remove_template_placeholders(slide):
        for shape in list(slide.shapes):
            # Remove existing placeholder text boxes or non-template overlays
            if shape.name in ["TextBox 8", "TextBox 9"] or (shape.has_text_frame and (
                "Problem Statement ID" in shape.text_frame.text or 
                "Proposed Solution (Describe your Idea" in shape.text_frame.text or 
                "Technologies to be used" in shape.text_frame.text or 
                "Analysis of the feasibility" in shape.text_frame.text or 
                "Potential impact" in shape.text_frame.text or 
                "Details / Links of the reference" in shape.text_frame.text
            )):
                sp = shape._element
                sp.getparent().remove(sp)
            # Remove any flat pictures added in content area (preserving template logos)
            elif shape.name.startswith("Picture") and shape.left > Inches(0.2) and shape.top > Inches(1.1):
                sp = shape._element
                sp.getparent().remove(sp)

    def add_card(slide, left, top, width, height, bg_color=C_WHITE, border_color=C_CARD_BORDER, border_width=1.5):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_color
        if border_color:
            shape.line.color.rgb = border_color
            shape.line.width = Pt(border_width)
        else:
            shape.line.fill.background()
        return shape

    def add_arrow(slide, left, top, width=0.25, height=0.20, color=C_BLUE_MED):
        arrow = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(left), Inches(top), Inches(width), Inches(height))
        arrow.fill.solid()
        arrow.fill.fore_color.rgb = color
        arrow.line.fill.background()
        return arrow

    def update_oval_team(slide, team_text="Team\nPraxis"):
        for shape in slide.shapes:
            if shape.has_text_frame and "Your" in shape.text_frame.text and "Team" in shape.text_frame.text:
                clear_textbox(shape)
                tf = shape.text_frame
                tf.vertical_anchor = MSO_ANCHOR.MIDDLE
                p = tf.paragraphs[0]
                p.alignment = PP_ALIGN.CENTER
                run = p.add_run()
                style_run(run, team_text, size=10.0, color=C_NAVY_DARK, bold=True)

    # =========================================================================
    # SLIDE 1: TITLE PAGE (100% Native Shapes)
    # =========================================================================
    s1 = prs.slides[0]
    remove_template_placeholders(s1)

    # Left Metadata Box
    s1_left = add_card(s1, 0.40, 1.35, 6.15, 4.35, bg_color=C_WHITE, border_color=C_BLUE_MED, border_width=1.8)
    ltf = s1_left.text_frame
    ltf.word_wrap = True
    ltf.margin_left = Inches(0.20)
    ltf.margin_right = Inches(0.20)
    ltf.margin_top = Inches(0.18)
    
    lp = ltf.paragraphs[0]
    lr = lp.add_run()
    style_run(lr, "PROBLEM STATEMENT & PROJECT DETAILS", size=13.0, color=C_BLUE_MED, bold=True)

    meta_items = [
        ("Problem Statement ID:", " SIH26160"),
        ("PS Title:", " AI-Powered IPsec VPN Protocol Analyzer & Security Assessment Framework"),
        ("Project Codename:", " IPsight (Solution v2.0)"),
        ("Sponsoring Agency:", " National Technical Research Organisation (NTRO)"),
        ("Theme & Category:", " Blockchain & Cybersecurity | Software"),
        ("Team ID & Name:", " 132859 | Team Praxis")
    ]
    for label, val in meta_items:
        p = ltf.add_paragraph()
        p.space_before = Pt(5)
        r1 = p.add_run()
        style_run(r1, f"• {label}", size=11.5, color=C_NAVY_DARK, bold=True)
        r2 = p.add_run()
        style_run(r2, val, size=11.5, color=C_TEXT_BODY, bold=(label in ["Problem Statement ID:", "Project Codename:", "Team ID & Name:"]))

    # Right Team Box
    s1_right = add_card(s1, 6.75, 1.35, 6.18, 4.35, bg_color=C_LIGHT_BG, border_color=C_BLUE_ACCENT, border_width=1.8)
    rtf = s1_right.text_frame
    rtf.word_wrap = True
    rtf.margin_left = Inches(0.20)
    rtf.margin_right = Inches(0.20)
    rtf.margin_top = Inches(0.18)
    
    rp = rtf.paragraphs[0]
    rr = rp.add_run()
    style_run(rr, "TEAM PRAXIS — IIT JODHPUR (AISHE: U-0395)", size=13.0, color=C_BLUE_ACCENT, bold=True)

    team_members = [
        ("Team Leader:", "Kumari Ankita", "B.Tech AI & Data Science (2nd Year)"),
        ("Team Member:", "Nandini Dhawan", "B.Tech AI & Data Science (2nd Year)"),
        ("Team Member:", "Chirag Jha", "B.Tech AI & Data Science (2nd Year)"),
        ("Team Member:", "Mayank Jangid", "B.Tech AI & Data Science (2nd Year)"),
        ("Team Member:", "Aayush", "B.Tech AI & Data Science (2nd Year)"),
        ("Team Member:", "Rudra Pratap Singh Chauhan", "B.Tech AI & Data Science (2nd Year)")
    ]
    for role, name, dept in team_members:
        p = rtf.add_paragraph()
        p.space_before = Pt(4)
        r1 = p.add_run()
        style_run(r1, f"• {role} ", size=11.0, color=C_NAVY_DARK, bold=True)
        r2 = p.add_run()
        style_run(r2, f"{name} ", size=11.0, color=C_BLUE_MED, bold=True)
        r3 = p.add_run()
        style_run(r3, f"({dept})", size=10.0, color=C_TEXT_MUTED, bold=False)

    # Bottom Core Principle Banner
    s1_banner = add_card(s1, 0.40, 5.85, 12.53, 0.95, bg_color=C_CYAN_BG, border_color=C_BLUE_ACCENT, border_width=1.5)
    btf = s1_banner.text_frame
    btf.vertical_anchor = MSO_ANCHOR.MIDDLE
    btf.margin_left = Inches(0.20)
    btf.margin_right = Inches(0.20)
    bp1 = btf.paragraphs[0]
    r_b1 = bp1.add_run()
    style_run(r_b1, "CORE ARCHITECTURAL PRINCIPLE: ", size=11.0, color=C_BLUE_MED, bold=True)
    r_b2 = bp1.add_run()
    style_run(r_b2, '"Directly observe what IPsec exposes; infer only what metadata supports; label every inference with confidence."', size=11.0, color=C_TEXT_DARK, bold=True, italic=True)
    
    bp2 = btf.add_paragraph()
    bp2.space_before = Pt(2)
    r_b3 = bp2.add_run()
    style_run(r_b3, "100% Offline & Air-Gapped  |  Zero-Key Passive Wire Inspection  |  Strict 3-Tier Evidence Provenance  |  Automated Remediation", size=10.0, color=C_TEXT_MUTED, bold=False)

    # =========================================================================
    # SLIDE 2: PROPOSED SOLUTION (100% Native Shapes)
    # =========================================================================
    s2 = prs.slides[1]
    update_oval_team(s2)
    remove_template_placeholders(s2)

    for shape in s2.shapes:
        if shape.has_text_frame and "IDEA TITLE" in shape.text_frame.text:
            clear_textbox(shape)
            tf = shape.text_frame
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            r = p.add_run()
            style_run(r, "IPsight — AI-Powered IPsec VPN Protocol Analyzer & Security Assessment Framework", size=15.0, color=C_NAVY_TITLE, bold=True)

    c_w = 4.00
    c_h = 5.45
    top_pos = 1.30

    # Card 1: Proposed Solution & Pipeline
    c1 = add_card(s2, 0.40, top_pos, c_w, c_h, bg_color=C_WHITE, border_color=C_BLUE_MED, border_width=1.8)
    tf1 = c1.text_frame
    tf1.word_wrap = True
    tf1.margin_left = Inches(0.20)
    tf1.margin_right = Inches(0.20)
    tf1.margin_top = Inches(0.18)
    
    p = tf1.paragraphs[0]
    r = p.add_run()
    style_run(r, "1. PROPOSED SOLUTION & PIPELINE", size=12.0, color=C_BLUE_MED, bold=True)
    
    p = tf1.add_paragraph()
    p.space_before = Pt(4)
    r = p.add_run()
    style_run(r, "Automated, 100% offline IPsec audit framework:", size=10.5, color=C_TEXT_DARK, bold=True)

    p1_items = [
        ("End-to-End Pipeline:", "Automates testbed ➔ capture ➔ deep parsing ➔ AI flow inference ➔ security assessment ➔ vendor remediation in one zero-touch chain."),
        ("Zero-Key Passive Wire Audit:", "Inspects IKE exchanges & ESP headers passively without requiring device credentials or private keys."),
        ("5-Head Calibrated AI Engine:", "Identifies IKE versions, tunnel/transport modes, crypto profiles, anomalous SPI churn, and in-tunnel traffic behavior."),
        ("Automated Multi-Vendor Fixes:", "Generates exact, copy-paste CLI remediation scripts for Cisco IOS-XE, Fortinet FortiOS, and strongSwan.")
    ]
    for k, v in p1_items:
        p = tf1.add_paragraph()
        p.space_before = Pt(5)
        r1 = p.add_run()
        style_run(r1, f"• {k} ", size=10.0, color=C_NAVY_DARK, bold=True)
        r2 = p.add_run()
        style_run(r2, v, size=10.0, color=C_TEXT_BODY, bold=False)

    # Card 2: How it Addresses the Problem
    c2 = add_card(s2, 4.66, top_pos, c_w, c_h, bg_color=C_WHITE, border_color=C_BLUE_ACCENT, border_width=1.8)
    tf2 = c2.text_frame
    tf2.word_wrap = True
    tf2.margin_left = Inches(0.20)
    tf2.margin_right = Inches(0.20)
    tf2.margin_top = Inches(0.18)
    
    p = tf2.paragraphs[0]
    r = p.add_run()
    style_run(r, "2. HOW IT ADDRESSES THE PROBLEM", size=12.0, color=C_BLUE_ACCENT, bold=True)

    p = tf2.add_paragraph()
    p.space_before = Pt(4)
    r = p.add_run()
    style_run(r, "Eliminating blind spots & audit bottlenecks:", size=10.5, color=C_TEXT_DARK, bold=True)

    p2_items = [
        ("Sub-2 Minute Speed:", "Replaces 3–5 days of manual Wireshark packet inspection with sub-2 minute automated audits (parses 100MB PCAP in <2 min)."),
        ("Detects Silent Downgrades:", "Flags when responders accept weak ciphers (3DES, SHA-1, DH <14) despite clients offering modern AEAD."),
        ("Exposes IKEv1 & PSK Leaks:", "Identifies deprecated IKEv1 Historic status and IKEv1 Aggressive Mode cleartext identity/hash exposures."),
        ("Encrypted Traffic Visibility:", "Infers in-tunnel traffic categories (VoIP, Web, Video, Bulk) with ZERO payload decryption, preserving total privacy.")
    ]
    for k, v in p2_items:
        p = tf2.add_paragraph()
        p.space_before = Pt(5)
        r1 = p.add_run()
        style_run(r1, f"• {k} ", size=10.0, color=C_NAVY_DARK, bold=True)
        r2 = p.add_run()
        style_run(r2, v, size=10.0, color=C_TEXT_BODY, bold=False)

    # Card 3: Innovation & Uniqueness
    c3 = add_card(s2, 8.92, top_pos, c_w, c_h, bg_color=C_WHITE, border_color=C_GREEN, border_width=1.8)
    tf3 = c3.text_frame
    tf3.word_wrap = True
    tf3.margin_left = Inches(0.20)
    tf3.margin_right = Inches(0.20)
    tf3.margin_top = Inches(0.18)
    
    p = tf3.paragraphs[0]
    r = p.add_run()
    style_run(r, "3. INNOVATION & UNIQUENESS", size=12.0, color=C_GREEN, bold=True)

    p = tf3.add_paragraph()
    p.space_before = Pt(4)
    r = p.add_run()
    style_run(r, "Defensible, defense-grade innovations:", size=10.5, color=C_TEXT_DARK, bold=True)

    p3_items = [
        ("3-Tier Evidence Taxonomy:", "Strictly categorizes parameters into DIRECT wire facts, INFERRED metadata models, and UNKNOWNS. Zero AI guessing."),
        ("Calibrated Behavior AI (SHAP):", "Uses FFT inter-arrival timing (VoIP ~20ms periodicity) with Platt/Isotonic calibration and Brier reliability scores."),
        ("Dual Decoupled Scoring:", "Calculates Security Posture (0–100, compliance) and Risk Urgency (0–100, threat level) + Likelihood × Impact Threat Matrix."),
        ("Post-Quantum (PQC) Runway:", "Evaluates classical DH groups against RFC 9370 ML-KEM pathways to defend against 'Harvest Now, Decrypt Later' threats.")
    ]
    for k, v in p3_items:
        p = tf3.add_paragraph()
        p.space_before = Pt(5)
        r1 = p.add_run()
        style_run(r1, f"• {k} ", size=10.0, color=C_NAVY_DARK, bold=True)
        r2 = p.add_run()
        style_run(r2, v, size=10.0, color=C_TEXT_BODY, bold=False)

    # =========================================================================
    # SLIDE 3: TECHNICAL APPROACH (100% Native Editable Flowchart & Cards)
    # =========================================================================
    s3 = prs.slides[2]
    update_oval_team(s3)
    remove_template_placeholders(s3)

    # Top Row: 4 Stage Cards (Native Shapes & Text)
    pipe_w = 2.80
    pipe_h = 2.70
    pipe_top = 1.30
    pipe_gap = 0.44

    pipeline_stages = [
        ("STAGE 1: INGEST", C_BLUE_MED, [
            ("Input Feeds:", "PCAP, Live Tap, Configs"),
            ("Stream Splitter:", "UDP 500/4500 (IKE) vs ESP"),
            ("Sessionizer:", "5-Tuple + SPI tracking"),
            ("Throughput:", "⚡ 100MB in <2 mins")
        ]),
        ("STAGE 2: DISSECT", C_PURPLE, [
            ("IKE Parser:", "IKEv1/v2 proposals & DH"),
            ("ESP Tracker:", "SPI integrity & jitter"),
            ("Vendor IDs:", "Fingerprints ➔ CVE lookup"),
            ("Data Model:", "⚡ Normalized JSON")
        ]),
        ("STAGE 3: AI INFER", C_AMBER, [
            ("H1/H2 Classifiers:", "Protocol & Mode"),
            ("H3 Crypto:", "Direct suite / ESP profile"),
            ("H4 Flow AI:", "FFT timing (VoIP ~20ms)"),
            ("H5 Outliers:", "⚡ Isolation Forest")
        ]),
        ("STAGE 4: ASSESS & FIX", C_GREEN, [
            ("Rule Engine:", "NIST SP 800-77 / RFC 8247"),
            ("Dual Scores:", "Posture & Risk (0-100)"),
            ("Fix Engine:", "1-Click Cisco, Fortinet, swanctl"),
            ("Output:", "⚡ PDF, SARIF & SHA-256")
        ])
    ]

    for idx, (st_title, st_color, st_items) in enumerate(pipeline_stages):
        l_pos = 0.40 + idx * (pipe_w + pipe_gap)
        
        # Native Card
        card = add_card(s3, l_pos, pipe_top, pipe_w, pipe_h, bg_color=C_WHITE, border_color=st_color, border_width=1.8)
        ctf = card.text_frame
        ctf.word_wrap = True
        ctf.margin_left = Inches(0.12)
        ctf.margin_right = Inches(0.12)
        ctf.margin_top = Inches(0.10)
        
        p = ctf.paragraphs[0]
        r = p.add_run()
        style_run(r, st_title, size=11.0, color=st_color, bold=True)
        
        for k, v in st_items:
            p = ctf.add_paragraph()
            p.space_before = Pt(3)
            r1 = p.add_run()
            style_run(r1, f"• {k} ", size=9.0, color=C_NAVY_DARK, bold=True)
            r2 = p.add_run()
            style_run(r2, v, size=9.0, color=(C_GREEN if v.startswith("⚡") else C_TEXT_BODY), bold=v.startswith("⚡"))

        # Add Native Arrow Connector between cards
        if idx < 3:
            arrow_left = l_pos + pipe_w + 0.08
            arrow_top = pipe_top + pipe_h / 2 - 0.10
            add_arrow(s3, arrow_left, arrow_top, width=0.28, height=0.20, color=C_BLUE_MED)

    # Bottom Left: Tech Stack Box (Native Card)
    s3_bot_top = 4.20
    b_left = add_card(s3, 0.40, s3_bot_top, 5.95, 2.55, bg_color=C_LIGHT_BG, border_color=C_BLUE_MED, border_width=1.8)
    btf1 = b_left.text_frame
    btf1.word_wrap = True
    btf1.margin_left = Inches(0.16)
    btf1.margin_right = Inches(0.16)
    btf1.margin_top = Inches(0.12)
    
    p = btf1.paragraphs[0]
    r = p.add_run()
    style_run(r, "TECHNOLOGIES & ARCHITECTURAL STACK", size=11.5, color=C_BLUE_MED, bold=True)
    
    tech_stack = [
        ("Packet Dissection:", " Python 3.11+, Scapy, tshark/dpkt, eBPF stream splitters"),
        ("AI / Machine Learning:", " scikit-learn, XGBoost, Isolation Forest, SHAP explainability"),
        ("Testbed Lab Factory:", " Docker-Compose, strongSwan 5.9+, iperf3, SIPp (VoIP), ffmpeg"),
        ("Backend & Remediation:", " FastAPI, SQLite / PostgreSQL, AST Multi-Vendor Fix Engine"),
        ("Frontend & Ecosystem:", " React 18 Web Dashboard + Expo Mobile SecOps Companion App")
    ]
    for k, v in tech_stack:
        p = btf1.add_paragraph()
        p.space_before = Pt(3)
        r1 = p.add_run()
        style_run(r1, f"• {k}", size=9.0, color=C_NAVY_DARK, bold=True)
        r2 = p.add_run()
        style_run(r2, v, size=9.0, color=C_TEXT_BODY, bold=False)

    # Bottom Right: 3-Tier Evidence Taxonomy Box (Native Card)
    b_right = add_card(s3, 6.55, s3_bot_top, 6.38, 2.55, bg_color=C_LIGHT_BG, border_color=C_GREEN, border_width=1.8)
    btf2 = b_right.text_frame
    btf2.word_wrap = True
    btf2.margin_left = Inches(0.16)
    btf2.margin_right = Inches(0.16)
    btf2.margin_top = Inches(0.12)
    
    p = btf2.paragraphs[0]
    r = p.add_run()
    style_run(r, "STRICT 3-TIER EVIDENCE PROVENANCE MODEL", size=11.5, color=C_GREEN, bold=True)

    evidence_rows = [
        ("✓ DIRECT (High Conf):", " IKE version, offered/accepted ciphers, DH group, Vendor ID, NAT-T — extracted verbatim from cleartext wire."),
        ("~ INFERRED (Calibrated):", " Flow category (VoIP, Web, Bulk, Video), ESP crypto profile — derived from timing & size distributions with zero decryption."),
        ("? UNKNOWN (Never Guess):", " Post-INIT IKEv2 secrets, CHILD_SA PFS without rekey, Raw PSK strings — strictly flagged to eliminate false claims.")
    ]
    for k, v in evidence_rows:
        p = btf2.add_paragraph()
        p.space_before = Pt(3)
        r1 = p.add_run()
        style_run(r1, f"{k}", size=9.0, color=C_NAVY_DARK, bold=True)
        r2 = p.add_run()
        style_run(r2, v, size=9.0, color=C_TEXT_BODY, bold=False)

    # =========================================================================
    # SLIDE 4: FEASIBILITY AND VIABILITY (100% Native Shapes)
    # =========================================================================
    s4 = prs.slides[3]
    update_oval_team(s4)
    remove_template_placeholders(s4)

    s4_top = 1.25
    s4_left = add_card(s4, 0.40, s4_top, 5.80, 5.50, bg_color=C_WHITE, border_color=C_BLUE_MED, border_width=1.8)
    s4_tf1 = s4_left.text_frame
    s4_tf1.word_wrap = True
    s4_tf1.margin_left = Inches(0.20)
    s4_tf1.margin_right = Inches(0.20)
    s4_tf1.margin_top = Inches(0.18)
    
    p = s4_tf1.paragraphs[0]
    r = p.add_run()
    style_run(r, "FEASIBILITY & PRACTICAL VIABILITY", size=13.0, color=C_BLUE_MED, bold=True)
    
    feas_points = [
        ("Lightweight & Low Footprint:", "Runs on standard commodity laptops or edge servers without requiring expensive GPU infrastructure."),
        ("High-Throughput Streaming:", "Processes large enterprise captures (<2 min for 100MB PCAP; <5 sec for router config files)."),
        ("100% Air-Gapped Operation:", "Completely self-contained with zero internet egress required; ideal for defense and classified environments."),
        ("Zero Performance Overhead:", "Operates 100% passively via network taps or span ports with zero latency impact on active VPN routers."),
        ("Multi-Vendor Interoperability:", "Standardized dissector engine supports Cisco, Fortinet, strongSwan, and cloud VPN gateways.")
    ]
    for k, v in feas_points:
        p = s4_tf1.add_paragraph()
        p.space_before = Pt(6)
        r1 = p.add_run()
        style_run(r1, f"• {k} ", size=10.5, color=C_NAVY_DARK, bold=True)
        r2 = p.add_run()
        style_run(r2, v, size=10.5, color=C_TEXT_BODY, bold=False)

    s4_right = add_card(s4, 6.40, s4_top, 6.53, 5.50, bg_color=C_WHITE, border_color=C_AMBER, border_width=1.8)
    s4_tf2 = s4_right.text_frame
    s4_tf2.word_wrap = True
    s4_tf2.margin_left = Inches(0.20)
    s4_tf2.margin_right = Inches(0.20)
    s4_tf2.margin_top = Inches(0.18)
    
    p = s4_tf2.paragraphs[0]
    r = p.add_run()
    style_run(r, "CHALLENGES & EXACT MITIGATIONS", size=13.0, color=C_AMBER, bold=True)

    challenges = [
        ("No Open Labeled Datasets:", "Our M1 Auto Testbed acts as a dataset factory, generating 12+ matrixed scenarios across ciphers, DH groups, and traffic types."),
        ("Encrypted ESP Opacity:", "Strict 3-Tier Taxonomy extracts IKE parameters directly and infers flow behaviors statistically without overclaiming cipher recovery."),
        ("Traffic Shaping & Jitter:", "Platt & Isotonic calibration outputs explicit Brier confidence scores, ensuring the score remains anchored on direct rules."),
        ("Untrusted PCAP Exploits:", "Parsers run inside sandboxed worker processes with memory and CPU bounds to prevent host compromise.")
    ]
    for ch, mit in challenges:
        p = s4_tf2.add_paragraph()
        p.space_before = Pt(6)
        r1 = p.add_run()
        style_run(r1, f"▲ {ch} ", size=10.5, color=C_RED, bold=True)
        r2 = p.add_run()
        style_run(r2, mit, size=10.5, color=C_TEXT_BODY, bold=False)

    # =========================================================================
    # SLIDE 5: IMPACT AND BENEFITS (100% Native Shapes)
    # =========================================================================
    s5 = prs.slides[4]
    update_oval_team(s5)
    remove_template_placeholders(s5)

    s5_top = 1.25
    s5_left = add_card(s5, 0.40, s5_top, 5.80, 5.50, bg_color=C_WHITE, border_color=C_BLUE_MED, border_width=1.8)
    s5_tf1 = s5_left.text_frame
    s5_tf1.word_wrap = True
    s5_tf1.margin_left = Inches(0.20)
    s5_tf1.margin_right = Inches(0.20)
    s5_tf1.margin_top = Inches(0.18)
    
    p = s5_tf1.paragraphs[0]
    r = p.add_run()
    style_run(r, "TARGET STAKEHOLDERS IMPACT", size=13.0, color=C_BLUE_MED, bold=True)

    stakeholders = [
        ("NTRO & Defense Commands:", "Enables continuous, passive wire-truth audit of sovereign military and strategic VPNs without needing device credentials."),
        ("CERT-In & Regulators:", "Automates national IPsec compliance inspections with standardized RFC/NIST citations and tamper-evident SHA-256 reports."),
        ("Critical Telecom Infrastructure:", "Audits thousands of enterprise WAN tunnels in seconds, identifying high-risk legacy suites across multi-tenant backbones."),
        ("Enterprise SecOps Teams:", "Provides real-time visibility into anomalous tunnels and delivers 1-click remediation scripts to close security gaps.")
    ]
    for k, v in stakeholders:
        p = s5_tf1.add_paragraph()
        p.space_before = Pt(6)
        r1 = p.add_run()
        style_run(r1, f"• {k} ", size=10.5, color=C_NAVY_DARK, bold=True)
        r2 = p.add_run()
        style_run(r2, v, size=10.5, color=C_TEXT_BODY, bold=False)

    s5_right = add_card(s5, 6.40, s5_top, 6.53, 5.50, bg_color=C_WHITE, border_color=C_GREEN, border_width=1.8)
    s5_tf2 = s5_right.text_frame
    s5_tf2.word_wrap = True
    s5_tf2.margin_left = Inches(0.20)
    s5_tf2.margin_right = Inches(0.20)
    s5_tf2.margin_top = Inches(0.18)
    
    p = s5_tf2.paragraphs[0]
    r = p.add_run()
    style_run(r, "QUANTIFIABLE BENEFITS & ROI", size=13.0, color=C_GREEN, bold=True)

    benefits = [
        ("90%+ Time & Cost Savings:", "Reduces audit cycles from 3–5 manual engineering days to <2 minutes per gateway, saving significant operational expenditure."),
        ("Zero-Decryption Privacy:", "Delivers behavioral intelligence (VoIP, Web, Video, Bulk) purely from metadata without decrypting citizen or enterprise payloads."),
        ("Elimination of Blind Spots:", "Detects hidden legacy ciphers (3DES, MD5, SHA-1, DH <14) and IKEv1 Aggressive Mode leaks before attackers exploit them."),
        ("Post-Quantum Readiness (PQC):", "Evaluates classical DH groups against RFC 9370 ML-KEM pathways, defending against 'Harvest Now, Decrypt Later' quantum attacks.")
    ]
    for k, v in benefits:
        p = s5_tf2.add_paragraph()
        p.space_before = Pt(6)
        r1 = p.add_run()
        style_run(r1, f"• {k} ", size=10.5, color=C_NAVY_DARK, bold=True)
        r2 = p.add_run()
        style_run(r2, v, size=10.5, color=C_TEXT_BODY, bold=False)

    # =========================================================================
    # SLIDE 6: RESEARCH AND REFERENCES (100% Native Shapes)
    # =========================================================================
    s6 = prs.slides[5]
    update_oval_team(s6)
    remove_template_placeholders(s6)

    s6_top = 1.25
    s6_left = add_card(s6, 0.40, s6_top, 5.80, 4.10, bg_color=C_WHITE, border_color=C_BLUE_MED, border_width=1.8)
    s6_tf1 = s6_left.text_frame
    s6_tf1.word_wrap = True
    s6_tf1.margin_left = Inches(0.20)
    s6_tf1.margin_right = Inches(0.20)
    s6_tf1.margin_top = Inches(0.14)
    
    p = s6_tf1.paragraphs[0]
    r = p.add_run()
    style_run(r, "AUTHORITATIVE STANDARDS & RFCS", size=12.0, color=C_BLUE_MED, bold=True)

    rfc_refs = [
        ("RFC 7296 / RFC 4301 / RFC 4303:", "IKEv2 Protocol, Security Architecture for IP, and Encapsulating Security Payload (ESP)."),
        ("RFC 8247 & RFC 8221:", "Current Cryptographic Requirements for IKEv2 and ESP (Deprecation of 3DES, DES, MD5, SHA-1)."),
        ("IETF draft-ietf-ipsecme-ikev1-to-historic:", "Formal transition of IKEv1 to Historic status."),
        ("NIST SP 800-77 Rev. 1 & NIST SP 800-131A:", "Mandates IKEv2, DH groups >=14, and AES-GCM AEAD ciphers."),
        ("RFC 9370 & NSA CNSA 2.0:", "Post-Quantum Hybrid Key Exchange (ML-KEM / Kyber transition).")
    ]
    for k, v in rfc_refs:
        p = s6_tf1.add_paragraph()
        p.space_before = Pt(4)
        r1 = p.add_run()
        style_run(r1, f"• {k} ", size=9.5, color=C_NAVY_DARK, bold=True)
        r2 = p.add_run()
        style_run(r2, v, size=9.5, color=C_TEXT_BODY, bold=False)

    s6_right = add_card(s6, 6.40, s6_top, 6.53, 4.10, bg_color=C_WHITE, border_color=C_GREEN, border_width=1.8)
    s6_tf2 = s6_right.text_frame
    s6_tf2.word_wrap = True
    s6_tf2.margin_left = Inches(0.20)
    s6_tf2.margin_right = Inches(0.20)
    s6_tf2.margin_top = Inches(0.14)
    
    p = s6_tf2.paragraphs[0]
    r = p.add_run()
    style_run(r, "ACADEMIC FOUNDATIONS & PRIOR ART", size=12.0, color=C_GREEN, bold=True)

    acad_refs = [
        ("Lundberg & Lee (SHAP Explainability):", "Feature attribution explaining model flow classifications (timing periodicity, burst symmetry)."),
        ("Liu et al. (Isolation Forest):", "High-dimensional outlier isolation for detecting proposal storms and sequence anomalies."),
        ("Guo et al. (Model Calibration):", "Isotonic regression & Platt scaling to guarantee well-calibrated confidence scores."),
        ("Wireshark & Zeek Benchmarks:", "Manual packet inspection tools; IPsight extends with automated scoring and 1-click multi-vendor fixes.")
    ]
    for k, v in acad_refs:
        p = s6_tf2.add_paragraph()
        p.space_before = Pt(4)
        r1 = p.add_run()
        style_run(r1, f"• {k} ", size=9.5, color=C_NAVY_DARK, bold=True)
        r2 = p.add_run()
        style_run(r2, v, size=9.5, color=C_TEXT_BODY, bold=False)

    s6_banner = add_card(s6, 0.40, 5.45, 12.53, 1.35, bg_color=C_CYAN_BG, border_color=C_BLUE_ACCENT, border_width=1.5)
    s6_btf = s6_banner.text_frame
    s6_btf.word_wrap = True
    s6_btf.margin_left = Inches(0.20)
    s6_btf.margin_right = Inches(0.20)
    s6_btf.margin_top = Inches(0.12)
    
    p = s6_btf.paragraphs[0]
    r1 = p.add_run()
    style_run(r1, "CORE NOVELTY CLAIM: ", size=10.5, color=C_BLUE_MED, bold=True)
    r2 = p.add_run()
    style_run(r2, "IPsight does not attempt to break tunnels. It inspects wire truth, detects downgrade risks, classifies encrypted flows with calibrated confidence, and auto-generates multi-vendor fixes.", size=10.0, color=C_TEXT_DARK, bold=True)

    p2 = s6_btf.add_paragraph()
    p2.space_before = Pt(3)
    r3 = p2.add_run()
    style_run(r3, "✓ M1 Auto Testbed Factory  |  ✓ 5-Head Calibrated AI Engine  |  ✓ Web Dashboard & Mobile App  |  ✓ Automated Cisco/Fortinet Fixes", size=9.5, color=C_GREEN, bold=True)

    # -------------------------------------------------------------------------
    # DELETE SLIDE 7 (Instruction Slide)
    # -------------------------------------------------------------------------
    if len(prs.slides) > 6:
        rId = prs.slides._sldIdLst[6].rId
        prs.part.drop_rel(rId)
        del prs.slides._sldIdLst[6]
        print(f"Slide 7 removed. Total slides remaining: {len(prs.slides)}")

    prs.save(output_pptx)
    prs.save(mirror_pptx)
    print(f"Presentation saved successfully to: {output_pptx} and {mirror_pptx}")

    # Re-export PDF
    try:
        import comtypes.client
        powerpoint = comtypes.client.CreateObject("Powerpoint.Application")
        powerpoint.Visible = 1
        deck = powerpoint.Presentations.Open(os.path.abspath(output_pptx))
        deck.SaveAs(os.path.abspath(output_pdf), 32)
        deck.SaveAs(os.path.abspath(mirror_pdf), 32)
        deck.Close()
        powerpoint.Quit()
        print(f"PDF saved successfully to: {output_pdf} and {mirror_pdf}")
    except Exception as e:
        print(f"PDF export notice: {e}")

if __name__ == "__main__":
    build_fully_editable_deck()
